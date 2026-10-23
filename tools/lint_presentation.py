#!/usr/bin/env python3
"""Rough layout linter for the generated deck.

It cannot replace looking at the slides, but it catches the two failures that are easy to
make in python-pptx: text that does not fit its box, and text boxes that overlap each other.

Run: python3 tools/lint_presentation.py
"""

from __future__ import annotations

import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "docs", "VR-Fire-Safety-Implementation.pptx")
FONT_PATH = os.path.join(ROOT, "Assets", "Fonts", "DejaVuSans.ttf")

EMU_PER_INCH = 914400


def parse_docx_line(line):
    return line


def measure(text, size_pt, width_pt, bold=False):
    """Return (line_count, height_pt) for wrapped text."""
    try:
        font = ImageFont.truetype(FONT_PATH, max(6, int(round(size_pt))))
    except OSError:
        return 1, size_pt * 1.3
    words = text.split(" ")
    lines = 1
    current = ""
    for word in words:
        trial = (current + " " + word).strip()
        if font.getlength(trial) <= width_pt or not current:
            current = trial
        else:
            lines += 1
            current = word
    return lines, lines * size_pt * 1.22


def main():
    prs = Presentation(DECK)
    slide_w = prs.slide_width
    slide_h = prs.slide_height
    problems = []

    for index, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            left, top = shape.left, shape.top
            width, height = shape.width, shape.height
            if left < 0 or top < 0 or left + width > slide_w + Emu(1000) \
                    or top + height > slide_h + Emu(1000):
                problems.append("slide %d: '%s' is outside the slide" % (index, shape.name))

            total_h = 0.0
            widest = 0
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs)
                if not text.strip():
                    total_h += 6
                    continue
                size = max((r.font.size.pt if r.font.size else 18) for r in para.runs)
                bold = any(r.font.bold for r in para.runs)
                width_pt = (width / EMU_PER_INCH) * 72
                lines, h = measure(text, size, width_pt, bold)
                line_spacing = para.line_spacing or 1.0
                total_h += h * line_spacing + (para.space_after.pt if para.space_after else 0)
                widest = max(widest, size)

            height_pt = (height / EMU_PER_INCH) * 72
            if total_h > height_pt + 4 and widest > 0:
                problems.append(
                    "slide %d: text needs %.0fpt but the box is %.0fpt — '%s'"
                    % (index, total_h, height_pt,
                       shape.text_frame.text.split("\n")[0][:52].replace("\n", " ")))

            if shape.text_frame.text.strip():
                boxes.append((shape.name, left, top, width, height,
                              shape.text_frame.text.strip().split("\n")[0][:40]))

        # overlap between text boxes (ignore deliberate band-on-panel stacking)
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ax1, ay1, ax2, ay2 = a[1], a[2], a[1] + a[3], a[2] + a[4]
                bx1, by1, bx2, by2 = b[1], b[2], b[1] + b[3], b[2] + b[4]
                ox = max(0, min(ax2, bx2) - max(ax1, bx1))
                oy = max(0, min(ay2, by2) - max(ay1, by1))
                if ox > 0 and oy > 0:
                    area = ox * oy
                    smaller = min(a[3] * a[4], b[3] * b[4])
                    if smaller and area / smaller > 0.35:
                        problems.append("slide %d: '%s' overlaps '%s' (%d%% of the smaller box)"
                                        % (index, a[5], b[5], 100 * area / smaller))

    print("Slides: %d" % len(prs.slides))
    if problems:
        print("\n%d layout warning(s):" % len(problems))
        for p in problems:
            print("  -", p)
        return 1
    print("No layout warnings.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
