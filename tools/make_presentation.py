#!/usr/bin/env python3
"""Builds the implementation slide deck for the VR Fire Safety trainer.

Run:  python3 tools/make_presentation.py
Out:  docs/VR-Fire-Safety-Implementation.pptx
"""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "VR-Fire-Safety-Implementation.pptx")

# ---------------------------------------------------------------- palette --
BG = RGBColor(0x0E, 0x15, 0x1E)
PANEL = RGBColor(0x17, 0x21, 0x2C)
PANEL_2 = RGBColor(0x1E, 0x2A, 0x37)
AMBER = RGBColor(0xF2, 0xA0, 0x3D)
RED = RGBColor(0xD9, 0x3A, 0x2B)
GREEN = RGBColor(0x3F, 0xBF, 0x7F)
BLUE = RGBColor(0x54, 0x9C, 0xD6)
TEXT = RGBColor(0xEC, 0xF1, 0xF6)
MUTED = RGBColor(0x9B, 0xA8, 0xB6)
DIM = RGBColor(0x6E, 0x7C, 0x8B)

FONT = "Segoe UI"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ utils --
def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    return s


def text(s, x, y, w, h, content, size=18, colour=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.1, italic=False, font=FONT, space_after=0):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = content.split("\n")
    for i, raw in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        # **bold** runs
        parts = raw.split("**")
        for j, part in enumerate(parts):
            if not part:
                continue
            run = p.add_run()
            run.text = part
            run.font.size = Pt(size)
            run.font.name = font
            run.font.bold = bold or (j % 2 == 1)
            run.font.italic = italic
            run.font.color.rgb = colour if (j % 2 == 0 or bold) else AMBER
    return box


def panel(s, x, y, w, h, fill=PANEL, line_colour=PANEL_2, radius=0.02, stripe=None):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_colour
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.adjustments[0] = radius
    if stripe:
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.02),
                                 Inches(0.07), Inches(h - 0.04))
        bar.fill.solid()
        bar.fill.fore_color.rgb = stripe
        bar.line.fill.background()
        bar.shadow.inherit = False
    return shape


def bullets(s, x, y, w, h, items, size=15, colour=TEXT, gap=7, bullet_colour=AMBER):
    """items: list of (level, text). level 0 = bullet, 1 = sub-bullet, 2 = plain line."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for level, raw in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.12
        p.space_after = Pt(gap)
        if level == 0:
            p.level = 0
            prefix = "▪  "
            size_use = size
        elif level == 1:
            p.level = 1
            prefix = "      –  "
            size_use = size - 1.5
        else:
            prefix = ""
            size_use = size
        run = p.add_run()
        run.text = prefix
        run.font.size = Pt(size_use)
        run.font.name = FONT
        run.font.color.rgb = bullet_colour if level < 2 else DIM
        parts = raw.split("**")
        for j, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size_use)
            r.font.name = FONT
            r.font.bold = j % 2 == 1
            r.font.color.rgb = (AMBER if j % 2 == 1 else colour)
    return box


def header(s, kicker, title, number=None):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.52), Inches(0.09),
                             Inches(0.26))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()
    bar.shadow.inherit = False
    text(s, 0.92, 0.5, 9.0, 0.3, kicker.upper(), size=11.5, colour=AMBER, bold=True)
    text(s, 0.7, 0.86, 11.6, 0.75, title, size=31, colour=TEXT, bold=True, line=1.0)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.68), Inches(1.3),
                              Inches(0.045))
    line.fill.solid()
    line.fill.fore_color.rgb = AMBER
    line.line.fill.background()
    line.shadow.inherit = False
    if number is not None:
        text(s, W - 1.3, 0.44, 0.6, 0.3, "%02d" % number, size=12, colour=DIM,
             align=PP_ALIGN.RIGHT)
    footer(s)


def footer(s):
    text(s, 0.7, H - 0.46, 8.0, 0.25, "VR Fire Safety & Evacuation Trainer  ·  implementation review",
         size=9, colour=DIM)
    text(s, W - 2.6, H - 0.46, 1.9, 0.25, "Unity 6000.5.3f1  ·  URP", size=9, colour=DIM,
         align=PP_ALIGN.RIGHT)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def chip(s, x, y, w, label, colour=AMBER, fill=PANEL_2):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w),
                               Inches(0.34))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = colour
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.adjustments[0] = 0.4
    tf = shape.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(11)
    r.font.name = FONT
    r.font.bold = True
    r.font.color.rgb = colour
    return shape


def arrow(s, x1, y1, x2, y2, colour=DIM, width=1.25):
    from pptx.enum.shapes import MSO_CONNECTOR
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                  Inches(x2), Inches(y2))
    conn.line.color.rgb = colour
    conn.line.width = Pt(width)
    return conn


# ------------------------------------------------------------------ 01 ------
s = slide()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.14))
band.fill.solid()
band.fill.fore_color.rgb = RED
band.line.fill.background()
band.shadow.inherit = False

text(s, 0.95, 1.15, 6.0, 0.3, "UNITY 6.5  ·  VR + DESKTOP  ·  ONE SCENE", size=12.5, colour=AMBER,
     bold=True)
text(s, 0.9, 1.6, 11.4, 1.9,
     "VR Fire Safety &\nEvacuation Trainer", size=52, colour=TEXT, bold=True, line=0.98)
text(s, 0.95, 3.55, 9.4, 0.9,
     "How the experience is built: the scene, the lesson engine, the interactions,\n"
     "the teaching content and the checks that keep it working.",
     size=17, colour=MUTED, line=1.25)

for i, (label, colour) in enumerate([
        ("Unity 6000.5.3f1", AMBER), ("URP 17.5", GREEN), ("OpenXR + XRI 3.5", BLUE),
        ("Desktop fallback", MUTED)]):
    chip(s, 0.95 + i * 2.35, 4.85, 2.2, label, colour)

line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(5.66), Inches(11.4),
                          Inches(0.012))
line.fill.solid()
line.fill.fore_color.rgb = PANEL_2
line.line.fill.background()
line.shadow.inherit = False

text(s, 0.95, 5.9, 6.2, 0.8,
     "**18 C# scripts  ·  4,300 lines**\n~270 scene objects, generated by a Python scene builder",
     size=13, colour=MUTED, line=1.35)
text(s, 7.4, 5.9, 5.0, 0.8,
     "**Educational simulation**\nNot a substitute for certified fire-safety training",
     size=13, colour=MUTED, line=1.35, align=PP_ALIGN.RIGHT)
notes(s, "Title slide. The deck explains how the trainer was implemented: environment, lesson "
         "engine, interaction system, teaching content, UI, and the automated verification. "
         "Everything shown is implemented in one Unity scene.")

# ------------------------------------------------------------------ 02 ------
s = slide()
header(s, "The brief", "What it teaches — and what it avoids", 2)

panel(s, 0.7, 1.95, 6.0, 4.55, stripe=AMBER)
text(s, 1.0, 2.15, 5.4, 0.4, "REQUIRED EXPERIENCE", size=12, colour=AMBER, bold=True)
bullets(s, 1.0, 2.62, 5.4, 3.7, [
    (0, "One **guided, linear** session for a beginner college audience"),
    (0, "Four teaching sections: alarm → masks/extinguishers → evacuation → first aid"),
    (0, "**One scene**: corridor, one stair flight, second corridor, assembly point"),
    (0, "Runs in the Editor **without a headset** (WASD / mouse / E) and in VR"),
    (0, "Real teaching content, cited from Indian and international guidance"),
    (0, "Large readable UI, no timers, no combat, no scoring"),
], size=14.5)

panel(s, 6.95, 1.95, 5.7, 4.55, stripe=RED)
text(s, 7.25, 2.15, 5.1, 0.4, "EXPLICITLY OUT OF SCOPE", size=12, colour=RED, bold=True)
bullets(s, 7.25, 2.62, 5.1, 3.7, [
    (0, "Extra floors, rooms, branches, campus, multiple incidents"),
    (0, "Separate mini-scenes per extinguisher type"),
    (0, "Scoring, leaderboards, AI, NPCs, multiplayer, architecture frameworks"),
    (0, "A **fake alternative route** past the smoke — teaches the honest alternative instead"),
    (0, "Claims of protection: no mask is presented as safe in smoke"),
    (0, "Claims of testing that was not actually performed"),
], size=14.5)

panel(s, 0.7, 6.5, 11.95, 0.44, fill=RGBColor(0x2A, 0x1B, 0x14), line_colour=RED)
text(s, 0.95, 6.59, 11.4, 0.3,
     "Presented as a basic educational simulation — not professional instruction or certification.",
     size=12, colour=AMBER)
notes(s, "Start from the constraints: the scope limits are what make a beginner project "
         "finishable. The out-of-scope list is deliberate and was respected throughout.")

# ------------------------------------------------------------------ 03 ------
s = slide()
header(s, "Environment", "One linear route through one building", 3)

stops = [
    ("START", "Corridor 1\nz 0 → 18 m", "Spawn point, tutorial signage"),
    ("SECTION A", "Alarm + phone\nleft / right wall", "Call point, 112 phone"),
    ("SECTION B", "Mask case +\n6 extinguishers", "Fire demo bin, PASS station"),
    ("LOBBY", "Stairs entrance\nsignage", "Lift call button (teaches: no lifts)"),
    ("STAIRS", "20 steps\ndown 3.5 m", "Handrails, EXIT signs"),
    ("SECTION C", "Corridor 2\nsmoke + fire door", "Hazard zone, safe point"),
    ("EXIT", "Double doors\nto outside", "Door must be opened"),
    ("SECTION D", "Assembly point\nz 51 m", "First-aid panels at the end"),
]
top = 2.15
step = 1.42
for i, (tag, title_, sub) in enumerate(stops):
    y = top + i * step * 0.0  # unused
    x = 0.75
    y = 2.05 + i * 0.6
    tag_colour = [AMBER, AMBER, AMBER, MUTED, BLUE, BLUE, GREEN, GREEN][i]
    pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.25),
                              Inches(0.44))
    pill.fill.solid()
    pill.fill.fore_color.rgb = PANEL_2
    pill.line.color.rgb = tag_colour
    pill.line.width = Pt(0.75)
    pill.shadow.inherit = False
    pill.adjustments[0] = 0.35
    tf = pill.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = tag
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = tag_colour

    text(s, 2.2, y - 0.02, 2.5, 0.5, title_, size=13.5, colour=TEXT, bold=True, line=1.05)
    text(s, 4.75, y + 0.04, 3.4, 0.5, sub, size=11.5, colour=MUTED, line=1.05)
    if i < len(stops) - 1:
        arrow(s, 1.37, y + 0.44, 1.37, y + 0.6 + 0.06, PANEL_2, 1.0)

panel(s, 8.55, 2.05, 4.1, 3.55, stripe=GREEN)
text(s, 8.85, 2.22, 3.6, 0.35, "HOW IT IS BUILT", size=11.5, colour=GREEN, bold=True)
bullets(s, 8.85, 2.62, 3.55, 2.9, [
    (0, "Primitive meshes + 53 generated materials"),
    (0, "Signs, audio and textures generated offline (Pillow / NumPy)"),
    (0, "Geometry written into the scene by **tools/gen_scene.py**"),
    (0, "No imported 3D assets, nothing to download"),
    (0, "Stairwell sealed top to bottom (ray-tested)"),
], size=12.5, gap=6)

panel(s, 8.55, 5.62, 4.1, 1.13, fill=PANEL_2)
text(s, 8.8, 5.78, 3.6, 0.85,
     "**Route length** ≈ 55 m\nCorridor 1 (y 0) → stairs (−3.5 m) → corridor 2 → outside",
     size=11.5, colour=MUTED, line=1.3)
notes(s, "Every teaching object sits on the single route, so the player never has to hunt or "
         "choose a branch. The stairs drop 3.5 m, the second corridor continues at that level, "
         "and the assembly point is outside.")

# ------------------------------------------------------------------ 04 ------
s = slide()
header(s, "Lesson design", "Four sections, one guided session", 4)

sections = [
    ("A", "Immediate response", "10 → 35",
     ["Interact with the alarm call point",
      "Alarm loop starts, instruction panel opens",
      "Call 112 on the wall phone",
      "Panel teaches: alert others, no belongings"]),
    ("B", "Masks & extinguishers", "40 → 85",
     ["Mask case → types, limits, rules",
      "Six models → class, rating, 'this fire?'",
      "Safe-response decision on a real bin fire",
      "Watch the animated PASS demonstration",
      "One small bin fire, one practice unit"]),
    ("C", "Evacuation exercise", "90 → 155",
     ["Alarm still sounding, route marked",
      "Lift call button → why not lifts",
      "Stairs → corridor 2 → smoke hazard",
      "Close the fire door, exit, assemble",
      "The smoke blocks the way - no shortcut"]),
    ("D", "First aid at the endpoint", "160 → 180",
     ["Burns: cool 20 min, cover, escalate",
      "Smoke inhalation: don't enter, call 112",
      "CPR only as directed by dispatch",
      "Restart or return to the menu",
      "Reading and decisions, no casualty drill"]),
]
for i, (letter, name, stages, items) in enumerate(sections):
    x = 0.7 + i * 3.05
    panel(s, x, 1.95, 2.85, 4.55, stripe=[AMBER, AMBER, BLUE, GREEN][i])
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.16), Inches(2.12), Inches(0.42),
                               Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = [AMBER, AMBER, BLUE, GREEN][i]
    badge.line.fill.background()
    badge.shadow.inherit = False
    tf = badge.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = letter
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = BG

    text(s, x + 0.66, 2.14, 2.05, 0.5, name, size=12.5, colour=TEXT, bold=True, line=1.0)
    text(s, x + 0.2, 2.76, 2.5, 0.3, "stages " + stages, size=10, colour=DIM)
    bullets(s, x + 0.2, 3.12, 2.5, 3.2, [(0, t) for t in items], size=12, gap=9)

text(s, 0.7, 6.5, 11.9, 0.4,
     "Progression is **event-driven**: each step unlocks the next interaction.", size=12,
     colour=MUTED)
notes(s, "The four sections map onto the brief. Section B is the longest because it carries the "
         "mask and extinguisher teaching; section C re-uses the alarm and the same building.")

# ------------------------------------------------------------------ 05 ------
s = slide()
header(s, "Architecture", "Small scripts around one lesson manager", 5)

# centre hub
hub = panel(s, 5.0, 3.02, 3.3, 1.05, fill=RGBColor(0x2C, 0x21, 0x12), line_colour=AMBER,
            stripe=None)
text(s, 5.12, 3.12, 3.1, 0.4, "LessonManager", size=17, colour=AMBER, bold=True,
     align=PP_ALIGN.CENTER)
text(s, 5.12, 3.52, 3.1, 0.4, "30 numbers = the whole lesson\nstage machine + event router",
     size=10.5, colour=MUTED, align=PP_ALIGN.CENTER, line=1.15)

left = [
    ("PlayerController", "WASD / mouse / XR sticks,\ninteraction raycast", 0.75, 1.95),
    ("AlarmController", "alarm loop, strobe,\nraise-once latch", 0.75, 2.95),
    ("SimpleInteractable", "eventId → LessonManager\n(phone, masks, lift, cards)", 0.75, 3.95),
    ("ExtinguisherStation", "animated PASS demo:\npin, aim, lever, sweep", 0.75, 4.95),
]
right = [
    ("LessonUI", "world-space UI, HUD,\noverlays, buttons", 8.95, 1.95),
    ("FireDemo", "bin fire, flames,\nlight, douse", 8.95, 2.95),
    ("EvacuationController", "checkpoints, smoke,\nfire door, exit", 8.95, 3.95),
    ("HudFollower", "panel follows the eyes,\nnever behind a wall", 8.95, 4.95),
]
for name, desc, x, y in left + right:
    panel(s, x, y, 3.55, 0.78, fill=PANEL_2)
    text(s, x + 0.16, y + 0.1, 3.25, 0.3, name, size=12.5, colour=TEXT, bold=True)
    text(s, x + 0.16, y + 0.38, 3.25, 0.35, desc, size=9.5, colour=MUTED, line=1.05)
    if x < 5:
        arrow(s, x + 3.55, y + 0.39, 5.0, 3.55, DIM, 1.25)
    else:
        arrow(s, 8.95, y + 0.39, 8.3, 3.55, DIM, 1.25)

panel(s, 0.75, 6.1, 5.7, 0.68, fill=PANEL_2)
text(s, 0.95, 6.2, 5.4, 0.5,
     "**LessonContent.cs** — all teaching text in one place (60 members)", size=11.5,
     colour=MUTED)
panel(s, 6.8, 6.1, 5.7, 0.68, fill=PANEL_2)
text(s, 7.0, 6.2, 5.4, 0.5,
     "**IInteractable / InteractableBase** — one contract for every usable object",
     size=11.5, colour=MUTED)
notes(s, "There is no framework and no DI container. Objects report events to LessonManager with "
         "a short string id (alarm, phone, mask_display, ext_card_water, pass_complete, "
         "checkpoint_stairs_bottom...). LessonManager decides what happens next.")

# ------------------------------------------------------------------ 06 ------
s = slide()
header(s, "Lesson engine", "A stage machine that reads like a script", 6)

code_items = [
    ("interaction", "alarm.Interact()", "→ Notify(\"alarm\")"),
    ("lesson", "stage == StageAlarm ? GoTo(StageAlarmInfo)", "shows the 112 panel"),
    ("unlock", "BeginPhoneStep()  → phone.available = true", "next object becomes usable"),
    ("gate", "everything else ignores the event", "no way to skip a section"),
]
y = 1.95
for i, (tag, a, b) in enumerate(code_items):
    panel(s, 0.7, y, 7.15, 0.72, fill=PANEL_2)
    text(s, 0.9, y + 0.08, 1.15, 0.3, tag.upper(), size=9.5, colour=AMBER, bold=True)
    text(s, 2.05, y + 0.06, 5.6, 0.3, a, size=12, colour=TEXT, font="Consolas")
    text(s, 2.05, y + 0.37, 5.6, 0.3, b, size=10.5, colour=MUTED, font="Consolas")
    y += 0.86

panel(s, 8.2, 1.95, 4.4, 3.3, stripe=BLUE)
text(s, 8.45, 2.1, 3.9, 0.35, "WHY THIS SHAPE", size=11.5, colour=BLUE, bold=True)
bullets(s, 8.45, 2.48, 3.9, 2.7, [
    (0, "One readable switch instead of scattered flags"),
    (0, "Every event is idempotent: a repeated trigger cannot skip ahead"),
    (0, "Reset is a single call — restart and menu both work"),
    (0, "New teaching step = one constant + one case"),
], size=11.5, gap=7)

stats = [("30", "stage constants"), ("17", "event ids handled"), ("5", "route checkpoints"),
         ("1", "scene, no reloads")]
for i, (big, small) in enumerate(stats):
    x = 0.7 + i * 2.95
    panel(s, x, 5.55, 2.75, 1.15, fill=PANEL_2)
    text(s, x + 0.2, 5.66, 1.2, 0.5, big, size=26, colour=AMBER, bold=True)
    text(s, x + 0.2, 6.16, 2.4, 0.3, small, size=10.5, colour=MUTED)
notes(s, "Stage constants are grouped in tens (10, 20 ... 180) with sub-steps in fives, so "
         "inserting a step later does not renumber the whole file. Unknown events are ignored, "
         "which keeps the lesson safe if an object is activated out of order.")

# ------------------------------------------------------------------ 07 ------
s = slide()
header(s, "Interaction", "Every object must answer — always", 7)

panel(s, 0.7, 1.95, 5.95, 2.5, stripe=AMBER)
text(s, 0.95, 2.1, 5.4, 0.35, "THE PIPELINE", size=11.5, colour=AMBER, bold=True)
bullets(s, 0.95, 2.48, 5.4, 1.9, [
    (0, "Camera ray + **SphereCastAll** (radius 0.22 m, 3.2 m reach)"),
    (0, "Hit collider → `GetComponentInParent<IInteractable>()`"),
    (0, "Skip the player's own colliders; nearest available wins"),
    (0, "HUD shows the prompt; **E** (desktop) or **trigger** (VR) activates"),
], size=12, gap=5)

panel(s, 0.7, 4.5, 5.95, 2.27, stripe=RED)
text(s, 0.95, 4.64, 5.4, 0.35, "THE BUG: SILENT OBJECTS", size=11.5, colour=RED, bold=True)
text(s, 0.95, 5.0, 5.45, 1.72,
     "Models and the mask case were locked until their exact lesson stage. Walk up early and "
     "pressing E did nothing — which reads as a broken game, not a lesson.\n\n"
     "**Fix:** they are always available. The stage only decides whether the interaction "
     "*counts*; otherwise it still opens the teaching panel.", size=12, colour=MUTED, line=1.2)

panel(s, 7.15, 1.95, 5.45, 4.82, stripe=GREEN)
text(s, 7.4, 2.1, 5.0, 0.35, "CONTRACT", size=11.5, colour=GREEN, bold=True)
code = [
    "interface IInteractable",
    "{",
    "    string InteractionPrompt { get; }",
    "    bool   IsAvailable       { get; }",
    "    void   Interact(PlayerController p);",
    "}",
]
for i, line in enumerate(code):
    text(s, 7.45, 2.5 + i * 0.32, 5.0, 0.3, line, size=11.5,
         colour=AMBER if "interface" in line else TEXT, font="Consolas")

text(s, 7.45, 4.5, 4.9, 0.35, "Derived classes:", size=11, colour=MUTED, bold=True)
bullets(s, 7.45, 4.85, 4.9, 1.8, [
    (0, "AlarmController · DoorController · ExtinguisherStation"),
    (0, "SimpleInteractable — one class, a string `eventId`, for phone, masks, lift and the six "
        "model cards"),
], size=11, gap=4)
notes(s, "The prompt, the hover pulse and the availability gate all live in InteractableBase, so "
         "no object can forget them. SimpleInteractable keeps the scene wiring to one serialized "
         "string per object.")

# ------------------------------------------------------------------ 08 ------
s = slide()
header(s, "Section A", "Raise the alarm, call 112, then leave", 8)

panel(s, 0.7, 1.95, 6.05, 2.35, stripe=AMBER)
text(s, 0.95, 2.1, 5.5, 0.35, "ALARM CALL POINT", size=11.5, colour=AMBER, bold=True)
bullets(s, 0.95, 2.48, 5.5, 1.75, [
    (0, "Red call point with strobe; `AlarmController` latches **raise-once**"),
    (0, "Two-tone hooter loop (700/950 Hz) starts, strobe pulses"),
    (0, "Panel then teaches: alert others, call 112, leave belongings"),
], size=12, gap=5)

panel(s, 0.7, 4.45, 6.05, 2.3, stripe=AMBER)
text(s, 0.95, 4.6, 5.5, 0.35, "112 PHONE", size=11.5, colour=AMBER, bold=True)
bullets(s, 0.95, 4.98, 5.5, 1.7, [
    (0, "Separate interactable on the opposite wall"),
    (0, "Teaches what to say: location, floor, what is burning, anyone trapped"),
    (0, "India's single emergency number — **112**, not 101/100"),
], size=12, gap=5)

panel(s, 7.25, 1.95, 5.35, 4.8, fill=RGBColor(0x22, 0x18, 0x12), line_colour=RED)
text(s, 7.5, 2.12, 4.85, 0.4, "TAUGHT IN THIS SECTION", size=11.5, colour=RED, bold=True)
bullets(s, 7.5, 2.55, 4.85, 4.0, [
    (0, "Raise the alarm **before** anything else"),
    (0, "Alert the people around you"),
    (0, "Call **112** and give a usable report"),
    (0, "Evacuate promptly, no belongings"),
    (0, "Do not go toward smoke or flames"),
    (0, "If you cannot leave: close a door, seal gaps, call, signal"),
], size=13.5, gap=13)
notes(s, "The alarm is a plain interactable: it starts audio and a strobe, and reports one event. "
         "It never takes control of the player, which was an early bug: an overlay that did not "
         "close left the player unable to move.")

# ------------------------------------------------------------------ 09 ------
s = slide()
header(s, "Section B · masks", "Teach at the object, not on a poster", 9)

masks = [
    ("Cloth / surgical", "Loose fit. No protection against CO or toxic fire gases.", RED),
    ("N95 / KN95", "Particles only, and only if sealed to a clean-shaven face.\nNo oxygen, no CO protection.", AMBER),
    ("Escape hood", "Filtered, short-duration escape aid. Needs oxygen in the air.\nNot a reason to enter smoke.", AMBER),
    ("SCBA", "Fire-service breathing apparatus with its own air supply.\nNot something to look for during an evacuation.", MUTED),
]
for i, (name, desc, colour) in enumerate(masks):
    y = 1.95 + i * 1.02
    panel(s, 0.7, y, 6.6, 0.9, fill=PANEL_2)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(y + 0.02), Inches(0.07),
                             Inches(0.86))
    bar.fill.solid()
    bar.fill.fore_color.rgb = colour
    bar.line.fill.background()
    bar.shadow.inherit = False
    text(s, 0.95, y + 0.1, 2.0, 0.3, name, size=13, colour=TEXT, bold=True)
    text(s, 3.05, y + 0.1, 4.1, 0.7, desc, size=10.5, colour=MUTED, line=1.15)

panel(s, 7.65, 1.95, 4.95, 2.1, stripe=GREEN)
text(s, 7.9, 2.1, 4.5, 0.35, "IMPLEMENTATION", size=11.5, colour=GREEN, bold=True)
bullets(s, 7.9, 2.48, 4.5, 1.5, [
    (0, "Mask case = `SimpleInteractable` (`mask_display`)"),
    (0, "Opens a panel with all four types + rules"),
    (0, "Reachable at any time, still counts once"),
], size=11.5, gap=5)

panel(s, 7.65, 4.2, 4.95, 2.55, fill=RGBColor(0x22, 0x18, 0x12), line_colour=RED)
text(s, 7.9, 4.35, 4.5, 0.35, "THE THREE RULES", size=11.5, colour=RED, bold=True)
bullets(s, 7.9, 4.72, 4.5, 1.9, [
    (0, "Never delay evacuation to look for a mask"),
    (0, "Never enter smoke because you wear one"),
    (0, "Staying low helps, but is **not** a guarantee"),
], size=12, gap=6)
notes(s, "The wall poster was removed on request: the same content is now delivered as a pop-up "
         "when the player inspects the case, where it is actually read.")

# ------------------------------------------------------------------ 10 ------
s = slide()
header(s, "Section B · extinguishers", "Six models, five classes, one question", 10)

panel(s, 0.7, 1.95, 6.1, 2.75, stripe=AMBER)
text(s, 0.95, 2.1, 5.5, 0.35, "INDIAN FIRE CLASSES", size=11.5, colour=AMBER, bold=True)
rows = [
    ("A", "solids: wood, paper, cloth", "water · foam · ABC"),
    ("B", "flammable liquids", "foam · ABC · CO₂"),
    ("C", "flammable gases", "ABC dry powder"),
    ("D", "combustible metals", "special Class D powder only"),
    ("F", "cooking oils and fats", "wet chemical"),
]
for i, (cls, what, used) in enumerate(rows):
    y = 2.45 + i * 0.42
    text(s, 0.98, y, 0.4, 0.3, cls, size=13, colour=AMBER, bold=True)
    text(s, 1.42, y + 0.02, 2.55, 0.3, what, size=11.5, colour=TEXT)
    text(s, 4.05, y + 0.02, 2.6, 0.3, used, size=11.5, colour=MUTED)
text(s, 0.98, 4.55, 5.6, 0.3,
     "Energised electrical equipment is a **hazard**, not a separate Indian class.",
     size=10.5, colour=MUTED, italic=True)

panel(s, 0.7, 4.85, 6.1, 1.9, stripe=RED)
text(s, 0.95, 4.98, 5.5, 0.35, "SAFETY RULES TAUGHT", size=11.5, colour=RED, bold=True)
bullets(s, 0.95, 5.35, 5.55, 1.35, [
    (0, "Never water on cooking oil or energised electrical equipment"),
    (0, "Never fight a leaking gas fire — alarm and evacuate"),
    (0, "Small, contained fires only; alarm raised; exit behind you"),
], size=11, gap=4)

panel(s, 7.1, 1.95, 5.5, 4.8, stripe=GREEN)
text(s, 7.35, 2.1, 5.0, 0.35, "CONTEXTUAL POP-UP AT EACH MODEL", size=11.5, colour=GREEN,
     bold=True)
text(s, 7.35, 2.5, 5.0, 0.35, "Every card returns three things:", size=11, colour=MUTED)
for i, (label, desc) in enumerate([
        ("Use it on", "the classes printed on the label"),
        ("This fire", "why it is / is not right for the bin"),
        ("Rating", "never above the label / one attempt only")]):
    y = 2.9 + i * 0.62
    tick = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.4), Inches(y + 0.02), Inches(0.22),
                              Inches(0.22))
    tick.fill.solid()
    tick.fill.fore_color.rgb = GREEN
    tick.line.fill.background()
    tick.shadow.inherit = False
    text(s, 7.75, y - 0.02, 1.3, 0.3, label, size=11.5, colour=TEXT, bold=True)
    text(s, 9.05, y - 0.02, 3.3, 0.3, desc, size=11, colour=MUTED)

text(s, 7.35, 4.9, 5.0, 0.35, "Example — the bin fire:", size=11, colour=MUTED, bold=True)
bullets(s, 7.35, 5.25, 5.0, 1.4, [
    (0, "**ABC powder** — right choice (Class A solids)"),
    (0, "**CO₂** — for liquids and live electrics; blows burning solids about"),
    (0, "**Wet chemical / Class D** — wrong fuel entirely"),
], size=10.5, gap=3)
notes(s, "One wall of models, one practice unit and one fire. The teaching text for all six units "
         "lives in LessonContent.Extinguishers and is bound to the models by event id.")

# ------------------------------------------------------------------ 11 ------
s = slide()
header(s, "Section B · PASS", "PASS in action — on the model itself", 11)

steps = [
    ("1", "PULL", "pin slides out of the valve,\nfalls and settles on the table"),
    ("2", "AIM", "valve, hose and horn rotate\nuntil the horn faces the base"),
    ("3", "SQUEEZE", "lever depresses;\npowder jet + hiss start"),
    ("4", "SWEEP", "head oscillates ±13°;\nflames and light fade out"),
]
for i, (num, title_, desc) in enumerate(steps):
    x = 0.7 + i * 3.08
    panel(s, x, 1.95, 2.9, 2.35, stripe=AMBER)
    text(s, x + 0.2, 2.05, 0.6, 0.45, num, size=24, colour=AMBER, bold=True)
    text(s, x + 0.72, 2.14, 2.0, 0.4, title_, size=15, colour=TEXT, bold=True)
    text(s, x + 0.2, 2.66, 2.55, 1.0, desc, size=10.5, colour=MUTED, line=1.2)
    if i < 3:
        arrow(s, x + 2.92, 3.1, x + 3.06, 3.1, AMBER, 1.5)

panel(s, 0.7, 4.55, 6.0, 2.2, stripe=BLUE)
text(s, 0.95, 4.68, 5.4, 0.35, "HOW IT IS IMPLEMENTED", size=11.5, colour=BLUE, bold=True)
bullets(s, 0.95, 5.05, 5.45, 1.6, [
    (0, "One interaction → coroutine sequence on the station"),
    (0, "Model has **Head / LeverPivot / Pin / SprayOrigin** transforms"),
    (0, "Spray = particle system parented to the horn, aimed along +Z"),
    (0, "`FireDemo.Douse()` fades flames, light and smoke on cue"),
], size=11.5, gap=4)

panel(s, 6.95, 4.55, 5.65, 2.2, stripe=GREEN)
text(s, 7.2, 4.68, 5.1, 0.35, "WHAT THE PLAYER SEES", size=11.5, colour=GREEN, bold=True)
bullets(s, 7.2, 5.05, 5.15, 1.6, [
    (0, "Large captions: “1 of 4  PULL the safety pin” …"),
    (0, "Real motion on the real model, in ~6 seconds"),
    (0, "Recap panel afterwards; station explains itself before"),
    (0, "Objective box tracks the step, no keys spammed"),
], size=11.5, gap=4)
notes(s, "This replaced a four-press interaction. The pin really leaves the valve and drops; the "
         "head really turns; the lever really travels; the jet really sweeps. All numbers are "
         "serialized on the station component so they can be tuned in the Inspector.")

# ------------------------------------------------------------------ 12 ------
s = slide()
header(s, "Section C", "The hazard is real; the shortcut is not", 12)

panel(s, 0.7, 1.95, 6.05, 2.15, stripe=BLUE)
text(s, 0.95, 2.1, 5.5, 0.35, "ROUTE CONTROL", size=11.5, colour=BLUE, bold=True)
bullets(s, 0.95, 2.48, 5.5, 1.55, [
    (0, "5 checkpoint triggers report their id to the lesson"),
    (0, "Lift call button → panel explains why lifts are not used"),
    (0, "Stairs: handrails both sides, EXIT signage, 20 steps"),
], size=12.5, gap=9)

panel(s, 0.7, 4.25, 6.05, 2.5, stripe=RED)
text(s, 0.95, 4.4, 5.5, 0.35, "SMOKE HAZARD", size=11.5, colour=RED, bold=True)
bullets(s, 0.95, 4.78, 5.5, 1.85, [
    (0, "`CorridorSmoke`: 4 particle systems, density ramps up"),
    (0, "`HazardZone` trigger blocks the dense part of the corridor"),
    (0, "Entering it resets the player to the safe point with a warning"),
    (0, "Fire door must be **closed** to protect the stairwell"),
], size=12.5, gap=9)

panel(s, 7.15, 1.95, 5.45, 2.5, fill=RGBColor(0x22, 0x18, 0x12), line_colour=AMBER)
text(s, 7.4, 2.1, 5.0, 0.35, "WHEN THE ROUTE IS BLOCKED", size=11.5, colour=AMBER, bold=True)
bullets(s, 7.4, 2.48, 5.0, 1.85, [
    (0, "Move to a **safer location**, not through the smoke"),
    (0, "**Close the door** if there is one"),
    (0, "Call **112** and describe where you are"),
    (0, "**Signal for help** — light, phone, window"),
], size=12, gap=7)

panel(s, 7.15, 4.6, 5.45, 2.15, stripe=GREEN)
text(s, 7.4, 4.73, 5.0, 0.35, "END OF THE ROUTE", size=11.5, colour=GREEN, bold=True)
bullets(s, 7.4, 5.1, 5.0, 1.55, [
    (0, "Exit doors open on interaction → outside"),
    (0, "Assembly marker + sign, then section D panels"),
    (0, "No timer, no failure state — only redirection"),
], size=12, gap=7)
notes(s, "The smoke is deliberately a hazard you cannot walk through: the trigger returns the "
         "player to the safe point and the panel teaches what to do instead of inventing a "
         "second corridor that does not exist.")

# ------------------------------------------------------------------ 13 ------
s = slide()
header(s, "Section D", "First aid: cautious, current, citable", 13)

panel(s, 0.7, 1.95, 6.0, 4.8, stripe=AMBER)
text(s, 0.95, 2.1, 5.4, 0.35, "BURNS", size=11.5, colour=AMBER, bold=True)
bullets(s, 0.95, 2.5, 5.45, 4.1, [
    (0, "Make the scene safe before helping"),
    (0, "Cool under **cool running water for 20 minutes**"),
    (0, "Remove jewellery and loose clothing — only if not stuck to the burn"),
    (0, "**No ice, toothpaste, oils, butter or turmeric**"),
    (0, "Do not burst blisters"),
    (0, "Cover loosely with a clean, non-fluffy dressing"),
    (0, "Urgent care for deep, large, facial, hand or airway burns"),
], size=13.5, gap=15)

panel(s, 6.95, 1.95, 5.65, 4.8, stripe=GREEN)
text(s, 7.2, 2.1, 5.1, 0.35, "SMOKE INHALATION", size=11.5, colour=GREEN, bold=True)
bullets(s, 7.2, 2.5, 5.15, 4.1, [
    (0, "**Never enter a dangerous area** to retrieve a casualty"),
    (0, "Move to fresh air only if it is safe to do so"),
    (0, "Call **112** early"),
    (0, "Monitor responsiveness and normal breathing"),
    (0, "Unresponsive and not breathing normally → follow the dispatcher's CPR instructions"),
    (0, "Use an AED if one is available and you are instructed to"),
], size=13.5, gap=15)
notes(s, "Delivered as panels at the safe endpoint — no casualty scenario, no new environment. "
         "Wording stays conservative: 'if safe', 'follow dispatch instructions'.")

# ------------------------------------------------------------------ 14 ------
s = slide()
header(s, "Interface", "One UI, two canvas modes, three inputs", 14)

panel(s, 0.7, 1.95, 5.9, 2.2, stripe=BLUE)
text(s, 0.95, 2.1, 5.3, 0.35, "CANVAS MODE IS CHOSEN AT RUNTIME", size=11.5, colour=BLUE,
     bold=True)
bullets(s, 0.95, 2.5, 5.35, 1.6, [
    (0, "**No headset** → Screen Space Overlay (drawn over the scene)"),
    (0, "**Headset** → World Space + `HudFollower` (overlay canvases do not render in VR)"),
    (0, "Switches automatically if a headset connects mid-session"),
], size=11, gap=4)

panel(s, 0.7, 4.3, 5.9, 2.45, stripe=AMBER)
text(s, 0.95, 4.45, 5.3, 0.35, "INPUT THAT CANNOT GET STUCK", size=11.5, colour=AMBER, bold=True)
bullets(s, 0.95, 4.83, 5.35, 1.85, [
    (0, "Buttons fire through `Button.onClick.Invoke()`, not only the EventSystem"),
    (0, "Mouse click · E / Enter / Space · arrows / Tab · VR trigger"),
    (0, "Self-heal: if the cursor stays locked while no overlay is open, control returns"),
], size=11, gap=5)

panel(s, 7.0, 1.95, 5.6, 4.8, stripe=GREEN)
text(s, 7.25, 2.1, 5.1, 0.35, "HUD, ALWAYS ON SCREEN", size=11.5, colour=GREEN, bold=True)
for i, (label, desc) in enumerate([
        ("Current objective", "top right, auto-shrinking text"),
        ("Section band", "top left, current of four"),
        ("Interaction prompt", "bottom centre, what E will do"),
        ("Caption band", "during the PASS demonstration"),
        ("Warning banner", "clear of the objective box"),
        ("Controls hint", "bottom left, always visible")]):
    y = 2.5 + i * 0.6
    text(s, 7.3, y, 2.3, 0.3, label, size=11.5, colour=TEXT, bold=True)
    text(s, 9.6, y, 2.8, 0.3, desc, size=11, colour=MUTED)

text(s, 7.25, 6.12, 5.15, 0.62,
     "1600 × 900 layout, scaled to the camera's field of view.",
     size=10.5, colour=DIM, italic=True)
notes(s, "The UI is built entirely in code (no prefabs), so it can never drift from the scripts. "
         "Two real bugs shaped it: buttons that only worked through the EventSystem, and a panel "
         "that opened inside a wall when the player stood at the alarm.")

# ------------------------------------------------------------------ 15 ------
s = slide()
header(s, "Content", "Checking the teaching, not just the code", 15)

rows = [
    ("Burns cooling — 20 minutes under cool running water", "WHO · NHS · Red Cross"),
    ("Masks do not stop CO; N95 filters particles only", "OSHA · 3M guidance"),
    ("Escape hoods are short-duration, oxygen-dependent", "OSHA interpretation"),
    ("Class C = flammable gases; A/B/C/D/F classes", "IS 15683 / IS 2190 · RDSO"),
    ("Single emergency number in India", "112.gov.in (MHA)"),
    ("CPR/AED only as dispatched / trained", "American Heart Association"),
]
for i, (claim, source) in enumerate(rows):
    y = 1.95 + i * 0.62
    panel(s, 0.7, y, 7.5, 0.54, fill=PANEL_2)
    text(s, 0.92, y + 0.11, 4.3, 0.3, claim, size=11, colour=TEXT)
    text(s, 5.3, y + 0.11, 2.8, 0.3, source, size=10.5, colour=AMBER, align=PP_ALIGN.RIGHT)

panel(s, 8.45, 1.95, 4.15, 2.55, stripe=GREEN)
text(s, 8.7, 2.1, 3.7, 0.35, "IN THE PROJECT", size=11.5, colour=GREEN, bold=True)
bullets(s, 8.7, 2.5, 3.7, 1.9, [
    (0, "`LessonContent.References` — titles + URLs"),
    (0, "Rendered on the in-world references board"),
    (0, "Only sources actually consulted"),
    (0, "No invented citations"),
], size=11, gap=5)

panel(s, 8.45, 4.65, 4.15, 2.1, fill=RGBColor(0x22, 0x18, 0x12), line_colour=RED)
text(s, 8.7, 4.78, 3.7, 0.35, "LANGUAGE RULES", size=11.5, colour=RED, bold=True)
bullets(s, 8.7, 5.15, 3.7, 1.5, [
    (0, "“if it is safe to do so”"),
    (0, "“follow the dispatcher's instructions”"),
    (0, "No absolutes, no heroics"),
], size=11, gap=4)
notes(s, "Accuracy was treated as a deliverable: each claim was checked against a named source "
         "and the references are shipped in the app itself.")

# ------------------------------------------------------------------ 16 ------
s = slide()
header(s, "Engineering", "Generated scene, automated checks", 16)

panel(s, 0.7, 1.95, 6.0, 2.3, stripe=AMBER)
text(s, 0.95, 2.1, 5.4, 0.35, "BUILD PIPELINE", size=11.5, colour=AMBER, bold=True)
bullets(s, 0.95, 2.5, 5.45, 1.6, [
    (0, "`gen_assets.py` → sign textures, WAV audio, 53 materials"),
    (0, "`gen_scene.py` → 276-object scene, meta files, build settings"),
    (0, "Deterministic: rerunning reproduces the same scene"),
], size=11.5, gap=5)

panel(s, 0.7, 4.4, 6.0, 2.35, stripe=BLUE)
text(s, 0.95, 4.55, 5.4, 0.35, "VERIFICATION INSTEAD OF GUESSING", size=11.5, colour=BLUE,
     bold=True)
bullets(s, 0.95, 4.93, 5.45, 1.7, [
    (0, "verify_scene: fileIDs, GUIDs, serialized fields, transforms, metas"),
    (0, "verify_flow: every event id handled, every stage reachable"),
    (0, "verify_project: editor version, packages, URP wiring, materials"),
    (0, "check_cs_access: cross-class member access (CS0122/CS0272 class)"),
], size=11, gap=4)

panel(s, 7.1, 1.95, 5.5, 4.8, stripe=RED)
text(s, 7.35, 2.1, 5.0, 0.35, "MIGRATED TO UNITY 6.5", size=11.5, colour=RED, bold=True)
bullets(s, 7.35, 2.5, 5.0, 4.1, [
    (0, "Target **6000.5.3f1**, Active Input Handling = Both, Linear"),
    (0, "Built-in RP → **URP 17.5**: pipeline asset, renderer, volume profile, "
        "53 materials, every light and the camera"),
    (0, "Removed APIs replaced (e.g. `FindFirstObjectByType`)"),
    (0, "Import-crash fix: ship the `Assets/XR` + `Assets/XRI/Settings` folders OpenXR "
        "expects, so it never calls `CreateFolder` on a missing path"),
    (0, "Sample-free OpenXR: enable the loader in XR Plug-in Management"),
], size=11.5, gap=10)
notes(s, "Nothing was accepted on faith: each generator change is followed by the four checks, "
         "and the URP migration is verified by inspecting the serialized assets, not by hoping "
         "the Editor upgrades them quietly.")

# ------------------------------------------------------------------ 17 ------
s = slide()
header(s, "Play-test findings", "What actually broke, and what fixed it", 17)

bugs = [
    ("Panels were invisible behind walls",
     "World-space UI sat 1.55 m ahead — inside the wall you face at the alarm. Desktop now uses "
     "Screen Space Overlay; VR pulls the canvas in front of geometry."),
    ("Menu text overlapped",
     "Bands collided at 16:9 and the canvas was cropped on other aspects. Rebuilt as stacked "
     "bands and scaled to the camera's field of view."),
    ("Sky visible after the stairs",
     "The stairwell is full height but the doorway is 2.3 m — the wall above it was missing. "
     "Sealed, plus a ray test and a scene check."),
    ("Extinguishers did nothing",
     "Models were locked until one lesson stage. Now always available; the stage only decides "
     "whether the interaction counts."),
    ("Particle console warning",
     "Smoke set only the X/Z velocity curves, leaving Y in another mode. All axes now match."),
]
for i, (title_, desc) in enumerate(bugs):
    y = 1.9 + i * 0.94
    panel(s, 0.7, y, 11.95, 0.82, fill=PANEL_2)
    num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.88), Inches(y + 0.2), Inches(0.4),
                             Inches(0.4))
    num.fill.solid()
    num.fill.fore_color.rgb = RED
    num.line.fill.background()
    num.shadow.inherit = False
    tf = num.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(i + 1)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = TEXT
    text(s, 1.45, y + 0.1, 3.5, 0.3, title_, size=12.5, colour=TEXT, bold=True)
    text(s, 5.05, y + 0.09, 7.4, 0.62, desc, size=10.5, colour=MUTED, line=1.15)

text(s, 0.7, 6.68, 11.9, 0.3,
     "Each finding was reproduced from a screenshot or the Console, fixed at the cause, then "
     "re-checked by the automated suite.", size=10.5, colour=DIM, italic=True)
notes(s, "This slide is the honest test log. Two of these came directly from the user's "
         "screenshots; the rest from static analysis of the generated scene.")

# ------------------------------------------------------------------ 18 ------
s = slide()
header(s, "Run it", "Open one scene and press Play", 18)

panel(s, 0.7, 1.95, 6.0, 2.85, stripe=AMBER)
text(s, 0.95, 2.1, 5.4, 0.35, "GETTING STARTED", size=11.5, colour=AMBER, bold=True)
steps = [
    "Unity Hub → Add → this folder → open with **6000.5.3f1**",
    "Open **Assets/Scenes/FireSafetyTrainer.unity**",
    "Press **Play** → main menu → Start Training",
    "VR: enable OpenXR in XR Plug-in Management, then Play",
]
for i, step in enumerate(steps):
    y = 2.55 + i * 0.55
    text(s, 0.98, y, 5.3, 0.3, "%d.  %s" % (i + 1, step), size=11.5, colour=TEXT, line=1.1)

panel(s, 0.7, 4.95, 6.0, 1.75, stripe=GREEN)
text(s, 0.95, 5.08, 5.4, 0.35, "DESKTOP CONTROLS", size=11.5, colour=GREEN, bold=True)
text(s, 0.95, 5.45, 5.45, 1.1,
     "**WASD** move   ·   **Mouse** look   ·   **Shift** run\n"
     "**E** interact   ·   **Click / Enter / Space** press buttons   ·   **Esc** pause",
     size=11.5, colour=MUTED, line=1.6)

panel(s, 7.05, 1.95, 5.55, 4.75, stripe=BLUE)
text(s, 7.3, 2.08, 5.0, 0.35, "WHAT IS VERIFIED — AND WHAT IS NOT", size=11.5, colour=BLUE,
     bold=True)

text(s, 7.3, 2.44, 5.05, 0.28, "VERIFIED", size=10, colour=GREEN, bold=True)
bullets(s, 7.3, 2.72, 5.08, 1.4, [
    (0, "Scene integrity, lesson flow, project settings, URP wiring, C# access, syntax"),
    (0, "Stairwell enclosure — ray test: no view escapes the building"),
    (0, "The project opens on 6000.5.3f1 with URP"),
], size=10.5, gap=4)

text(s, 7.3, 4.2, 5.05, 0.28, "NOT VERIFIED HERE", size=10, colour=RED, bold=True)
bullets(s, 7.3, 4.48, 5.08, 1.1, [
    (0, "VR runtime on a headset — no headset in this environment"),
    (0, "The rendered look of the new UI and the PASS animation"),
], size=10.5, gap=4, bullet_colour=RED)

text(s, 7.3, 5.75, 5.05, 0.55,
     "Manual step: enable OpenXR, then import the device simulator for headset preview.",
     size=9.5, colour=DIM, italic=True, line=1.15)
notes(s, "Close with the run instructions and an honest test statement. The deck does not claim "
         "any VR run was performed.")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("Wrote", OUT, "with", len(prs.slides), "slides")
