#!/usr/bin/env python3
"""Build the Word version of the project report from docs/Project-Report.md.

The Markdown file is the single source of truth: this script only lays it out for Word
(headings, tables, screenshots with captions, code block, page numbers). Edit the Markdown,
run `python3 tools/make_report_docx.py`, and the .docx follows.

    python3 tools/make_report_docx.py        # docs/VR-Fire-Safety-Project-Report.docx

Requires python-docx and Pillow (both pip-installable). The Unity project does not need
either: these are development tools.
"""

import os
import re
import sys
from PIL import Image

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
SOURCE = os.path.join(DOCS, "Project-Report.md")
TARGET = os.path.join(DOCS, "VR-Fire-Safety-Project-Report.docx")

INK = RGBColor(0x1B, 0x22, 0x2B)
MUTED = RGBColor(0x5B, 0x66, 0x72)
AMBER = RGBColor(0xA9, 0x64, 0x14)
RULE = "D8DEE5"

MAX_IMAGE_WIDTH = Inches(6.25)


# --------------------------------------------------------------------------- #
# Reading the Markdown
# --------------------------------------------------------------------------- #

def parse_inline(text):
    """Split a line into (text, bold, code) runs."""
    runs = []
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def is_table_row(line):
    return line.startswith("|") and line.rstrip().endswith("|")


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse(markdown):
    """Turn the report's Markdown subset into a list of blocks."""
    lines = markdown.split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", body))
            continue

        # horizontal rule
        if stripped == "---":
            blocks.append(("hr", None))
            i += 1
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2))))
            i += 1
            continue

        # image, then optional italic caption (which may run over several lines)
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if m:
            path = m.group(2)
            caption = []
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines):
                start = lines[j].strip()
                if start.startswith("*") and not start.startswith("**"):
                    while j < len(lines) and lines[j].strip():
                        text = lines[j].strip()
                        body = text[1:] if (text.startswith("*") and not text.startswith("**")) else text
                        closed = body.endswith("*")
                        if closed:
                            body = body[:-1]
                        if body.strip():
                            caption.append(body.strip())
                        j += 1
                        if closed:
                            break
            blocks.append(("img", (path, " ".join(caption))))
            i = j
            continue

        # tables
        if is_table_row(line):
            rows = []
            while i < len(lines) and is_table_row(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            blocks.append(("table", rows))
            continue

        # ordered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            items = []
            while i < len(lines):
                mm = re.match(r"^(\d+)\.\s+(.*)$", lines[i].strip())
                if mm:
                    items.append([mm.group(2)])
                    i += 1
                elif lines[i].startswith("   ") and lines[i].strip() and items:
                    items[-1].append(lines[i].strip())
                    i += 1
                else:
                    break
            blocks.append(("ol", items))
            continue

        # bullets
        if stripped.startswith("- "):
            items = []
            while i < len(lines):
                current = lines[i]
                if current.strip().startswith("- "):
                    items.append([current.strip()[2:]])
                    i += 1
                elif current.startswith("  ") and current.strip() and items:
                    items[-1].append(current.strip())
                    i += 1
                else:
                    break
            blocks.append(("ul", items))
            continue

        # paragraph
        paragraph = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not (
            lines[i].strip().startswith(("#", "- ", "```", "|", "!"))
            or re.match(r"^\d+\.\s+", lines[i].strip())
            or lines[i].strip() == "---"
        ):
            paragraph.append(lines[i].strip())
            i += 1
        blocks.append(("p", " ".join(paragraph)))

    return blocks


# --------------------------------------------------------------------------- #
# Word helpers
# --------------------------------------------------------------------------- #

def shade(cell, hex_fill):
    props = OxmlElement("w:shd")
    props.set(qn("w:val"), "clear")
    props.set(qn("w:color"), "auto")
    props.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(props)


def set_cell_margins(table, top=60, bottom=60, left=110, right=110):
    for row in table.rows:
        for cell in row.cells:
            mar = OxmlElement("w:tcMar")
            for tag, value in (("top", top), ("bottom", bottom),
                               ("start", left), ("end", right)):
                node = OxmlElement("w:" + tag)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")
                mar.append(node)
            cell._tc.get_or_add_tcPr().append(mar)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


def add_runs(paragraph, text, size=None, colour=None):
    for chunk, bold, code in parse_inline(text):
        run = paragraph.add_run(chunk)
        run.bold = bold
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 0.8)
            run.font.color.rgb = RGBColor(0x9A, 0x3B, 0x1E)
        else:
            if size:
                run.font.size = Pt(size)
            if colour is not None:
                run.font.color.rgb = colour


def rule_below(paragraph, colour=RULE, size=6):
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), colour)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


# --------------------------------------------------------------------------- #
# Building the document
# --------------------------------------------------------------------------- #

def build():
    with open(SOURCE, encoding="utf-8") as handle:
        blocks = parse(handle.read())

    document = Document()

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    headings = {1: (23, AMBER, 6), 2: (20, INK, 12), 3: (14, INK, 10), 4: (12, INK, 8)}
    for level, (size, colour, before) in headings.items():
        style = document.styles["Heading %d" % level]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = colour
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    # Header and footer
    header = section.header.paragraphs[0]
    header.text = "VR Fire Safety and Evacuation Trainer  ·  project report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    title_written = False
    subtitle_written = False
    first_section_seen = False
    for kind, payload in blocks:
        if kind == "h":
            level, text = payload
            if level == 1 and not title_written:
                title_written = True
                title = document.add_paragraph()
                title.paragraph_format.space_after = Pt(2)
                run = title.add_run(text)
                run.font.size = Pt(26)
                run.font.bold = True
                run.font.color.rgb = INK
                continue
            # A cover page, then a fresh page where the history section starts.
            if level == 2 and (not first_section_seen or text.startswith("13.")):
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if level == 2:
                first_section_seen = True
            heading = document.add_heading(level=min(level, 3))
            add_runs(heading, text)
            if level == 2:
                rule_below(heading, "E8EBEF", 4)
            continue

        if kind == "p":
            text = payload
            # The line right under the title is the subtitle.
            if title_written and not subtitle_written and text == "Project report":
                subtitle_written = True
                para = document.add_paragraph()
                para.paragraph_format.space_after = Pt(10)
                run = para.add_run("Project report")
                run.font.size = Pt(13)
                run.font.color.rgb = AMBER
                run.bold = True
                rule_below(para, "E3C79B", 8)
                continue
            para = document.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(para, text)
            continue

        if kind == "ul":
            for item in payload:
                para = document.add_paragraph(style="List Bullet")
                para.paragraph_format.space_after = Pt(3)
                add_runs(para, " ".join(item))
            continue

        if kind == "ol":
            for number, item in enumerate(payload, start=1):
                para = document.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.75)
                para.paragraph_format.first_line_indent = Cm(-0.75)
                para.paragraph_format.space_after = Pt(3)
                para.add_run("%d.  " % number).bold = True
                add_runs(para, " ".join(item))
            continue

        if kind == "table":
            rows = payload
            if not rows:
                continue
            # A table whose first row is blank has no header: drop it rather than print an
            # empty grey band above the data.
            has_header = any(cell.strip() for cell in rows[0])
            if not has_header:
                rows = rows[1:]
            columns = max(len(r) for r in rows)
            table = document.add_table(rows=0, cols=columns)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for column in range(columns):
                    text = row[column] if column < len(row) else ""
                    cell = cells[column]
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.space_before = Pt(1)
                    para.paragraph_format.line_spacing = 1.0
                    if index == 0 and has_header:
                        shade(cell, "EDEFF2")
                        add_runs(para, text, size=9.5)
                        for run in para.runs:
                            run.bold = True
                            run.font.color.rgb = INK
                    else:
                        add_runs(para, text, size=9.5)
            set_cell_margins(table)
            # Fill the text column, repeat the header row on a new page, and never split a
            # row across two pages.
            width = OxmlElement("w:tblW")
            width.set(qn("w:w"), "5000")
            width.set(qn("w:type"), "pct")
            table._tbl.tblPr.append(width)
            for index, row in enumerate(table.rows):
                props = row._tr.get_or_add_trPr()
                cant_split = OxmlElement("w:cantSplit")
                props.append(cant_split)
                if index == 0 and has_header:
                    header = OxmlElement("w:tblHeader")
                    props.append(header)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if kind == "code":
            for line in payload:
                para = document.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.left_indent = Cm(0.4)
                para.paragraph_format.line_spacing = 1.0
                run = para.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8.6)
                run.font.color.rgb = RGBColor(0x2E, 0x3A, 0x46)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if kind == "img":
            path, caption = payload
            full = os.path.join(DOCS, path)
            if not os.path.exists(full):
                print("  ! missing image:", full, file=sys.stderr)
                continue
            width, height = Image.open(full).size
            width_in = min(MAX_IMAGE_WIDTH, Inches(width / 150.0))
            para = document.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(full, width=width_in)
            if caption:
                cap = document.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(12)
                run = cap.add_run(caption)
                run.italic = True
                run.font.size = Pt(8.8)
                run.font.color.rgb = MUTED
            continue

        if kind == "hr":
            para = document.add_paragraph()
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            rule_below(para, RULE, 6)
            continue

    document.save(TARGET)
    print("Wrote", TARGET)


if __name__ == "__main__":
    build()
